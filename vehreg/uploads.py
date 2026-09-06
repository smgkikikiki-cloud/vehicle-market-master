"""Taking a file the owner has in a folder and turning it into loadable months.

The pieces to read every shape DLT publishes already existed — header sniffing
for CSV exports, the pivot and long readers for the workbooks, the resolver and
the review queue for labels. What did not exist was a way to hand the app a
file: the only route in was to copy it into ``data/raw`` and restart.

Everything here is Streamlit-free so it can be tested without a browser, and
nothing here writes to the live warehouse. A caller inspects a file, asks for a
dry run against a scratch database, and only then decides to keep it.
"""

from __future__ import annotations

import shutil
import sqlite3
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from . import dlt, dlt_pivot
from .catalog import DATA_DIR, Catalog, available_years
from .db import connect, rebuild_dimension
from .ingest import ColumnMap, ingest_csv, read_rows, sniff_columns

#: The shape every upload is normalised to before it is loaded. It is the one
#: the pivot importer already writes, so an uploaded month and a month imported
#: from a workbook are the same kind of file afterwards.
CANONICAL_HEADER = dlt_pivot.CSV_HEADER

WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass
class Upload:
    """What a file turned out to be."""

    path: Path
    kind: str                      # "workbook-long" | "workbook-pivot" | "csv"
    periods: dict[str, float]      # period -> units the file says it holds
    detail: str = ""               # how it was read, for the reader
    warnings: list[str] = field(default_factory=list)
    _rows: list = field(default_factory=list, repr=False)
    _colmap: ColumnMap | None = field(default=None, repr=False)

    @property
    def total(self) -> float:
        return sum(self.periods.values())


def inspect(path: Path | str) -> Upload:
    """Work out what the file is, without loading anything."""
    path = Path(path)
    if path.suffix.lower() in WORKBOOK_SUFFIXES:
        return _inspect_workbook(path)
    return _inspect_csv(path)


def _inspect_workbook(path: Path) -> Upload:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    names = book.sheetnames
    book.close()

    # The long sheet is worth more: it carries the registration class, without
    # which a pickup cannot be split by cab and a fifth of a month stops at the
    # brand. So it is tried first even when both sheets are present.
    for sheet in names:
        rows = dlt_pivot.load_workbook_rows(path, sheet)
        try:
            totals = dlt_pivot.long_period_totals(rows)
        except ValueError:
            continue
        if totals:
            return Upload(path=path, kind="workbook-long", periods=totals,
                          detail=f"ชีต {sheet!r} — มีคอลัมน์ประเภทรถ (รย.)",
                          _rows=rows)

    for sheet in names:
        rows = dlt_pivot.load_workbook_rows(path, sheet)
        try:
            totals = dlt_pivot.period_totals(rows)
            declared = dlt_pivot.declared_period_totals(rows)
        except ValueError:
            continue
        if not totals:
            continue
        warnings = []
        drift = [p for p in declared
                 if abs(totals.get(p, 0.0) - declared[p]) > 0.5]
        if drift:
            warnings.append(
                "ยอดรายรุ่นไม่ตรงกับยอดรวมที่ไฟล์เขียนไว้เอง ในเดือน "
                + ", ".join(sorted(drift)[:4])
            )
        return Upload(path=path, kind="workbook-pivot", periods=totals,
                      detail=f"ชีต {sheet!r} — ไม่มีคอลัมน์ประเภทรถ กระบะจะ"
                             "แยกตัวถังไม่ได้",
                      warnings=warnings, _rows=rows)

    raise ValueError(f"{path.name}: อ่านไม่ออกว่าเป็นชีตแบบไหน "
                     f"(มีชีต: {', '.join(names)})")


def _inspect_csv(path: Path) -> Upload:
    colmap, rows = read_rows(path)
    missing = colmap.missing()
    if missing:
        raise ValueError(
            f"{path.name}: หาคอลัมน์ไม่เจอ ({', '.join(missing)}) "
            "— ระบุเองในหน้าถัดไปหรือแก้หัวตารางในไฟล์"
        )
    totals: dict[str, float] = {}
    for row in rows:
        period = dlt.period_of(row.get(colmap.period)) if hasattr(dlt, "period_of") \
            else str(row.get(colmap.period) or "").strip()[:7]
        units = _number(row.get(colmap.units))
        if period and units:
            totals[period] = totals.get(period, 0.0) + units
    return Upload(path=path, kind="csv", periods=dict(sorted(totals.items())),
                  detail="อ่านหัวตารางได้: " + _describe(colmap),
                  _rows=rows, _colmap=colmap)


def _describe(colmap: ColumnMap) -> str:
    named = {"เดือน": colmap.period, "ยี่ห้อ": colmap.brand,
             "รุ่น": colmap.model, "จำนวน": colmap.units,
             "ประเภทรถ": colmap.registration_type}
    return " · ".join(f"{k}={v}" for k, v in named.items() if v)


def _number(raw) -> float:
    text = str(raw or "").replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def write_period(upload: Upload, period: str, target: Path | str) -> int:
    """Write one month of the upload in the canonical shape."""
    target = Path(target)
    if upload.kind == "workbook-long":
        return dlt_pivot.write_long_period_csv(upload._rows, period, target)
    if upload.kind == "workbook-pivot":
        return dlt_pivot.write_period_csv(upload._rows, period, target)
    return _write_csv_period(upload, period, target)


def _write_csv_period(upload: Upload, period: str, target: Path) -> int:
    import csv as _csv

    colmap = upload._colmap
    assert colmap is not None
    merged: dict[tuple[str, str, str], float] = {}
    for row in upload._rows:
        row_period = str(row.get(colmap.period) or "").strip()[:7]
        if row_period != period:
            continue
        units = _number(row.get(colmap.units))
        if not units:
            continue
        reg = (str(row.get(colmap.registration_type) or "").strip()
               if colmap.registration_type else "")
        key = (dlt_pivot.registration_code(reg) or dlt_pivot.ANY_CLASS,
               str(row.get(colmap.brand) or "").strip() if colmap.brand else "",
               str(row.get(colmap.model) or "").strip() if colmap.model else "")
        merged[key] = merged.get(key, 0.0) + units

    target.parent.mkdir(parents=True, exist_ok=True)
    with open(target, "w", newline="", encoding="utf-8-sig") as handle:
        writer = _csv.writer(handle)
        writer.writerow(CANONICAL_HEADER)
        for (reg, brand, model), units in sorted(merged.items()):
            writer.writerow([period, reg, brand, model, f"{units:g}"])
    return len(merged)


@dataclass
class DryRun:
    period: str
    rows: int
    units_in_file: float
    units_loaded: float
    units_review: float
    units_brand_grain: float
    already_loaded: float | None
    top_unmatched: list[tuple[str, float]]

    @property
    def brand_grain_share(self) -> float:
        return (self.units_brand_grain / self.units_loaded
                if self.units_loaded else 0.0)


def dry_run(upload: Upload, period: str, *, data_dir: Path | str = DATA_DIR,
            live_db: Path | str | None = None) -> DryRun:
    """Load one month into a scratch database and report what would happen.

    Nothing about the live warehouse changes. The point is to answer "how much
    of this file will actually classify" before the owner commits to it, which
    is the question the review queue only answers after the fact.
    """
    data_dir = Path(data_dir)
    year = int(period[:4])
    if year not in set(available_years(data_dir)):
        raise ValueError(f"ไม่มี catalog ปี {year} — โหลดเข้าไปทุกแถวจะตกไป review")

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        source = tmp / f"upload_{period}.csv"
        rows = write_period(upload, period, source)

        conn = connect(tmp / "scratch.sqlite3")
        catalog = Catalog.load(data_dir, year)
        rebuild_dimension(conn, catalog)
        ingest_csv(conn, catalog, source, f"UPLOAD {period}",
                   allow_duplicate_payload=True)

        loaded = conn.execute(
            "SELECT COALESCE(SUM(units),0) FROM fact_registration"
        ).fetchone()[0]
        brand_grain = conn.execute(
            "SELECT COALESCE(SUM(f.units),0) FROM fact_registration f "
            "JOIN dim_unit u ON u.unit_id = f.unit_id "
            "AND u.catalog_year = ? WHERE u.grain = 'BRAND'", (year,)
        ).fetchone()[0]
        review = conn.execute(
            "SELECT COALESCE(SUM(units),0) FROM ingest_review WHERE status='open'"
        ).fetchone()[0]
        top = conn.execute(
            "SELECT raw_label, SUM(units) AS u FROM ingest_review "
            "WHERE status='open' GROUP BY 1 ORDER BY u DESC LIMIT 8"
        ).fetchall()
        conn.close()

    already = None
    if live_db is not None and Path(live_db).exists():
        live = connect(Path(live_db))
        row = live.execute(
            "SELECT COALESCE(SUM(units),0) FROM fact_registration WHERE period=?",
            (period,)).fetchone()
        already = float(row[0]) if row else 0.0
        live.close()

    return DryRun(
        period=period, rows=rows, units_in_file=upload.periods.get(period, 0.0),
        units_loaded=float(loaded), units_review=float(review),
        units_brand_grain=float(brand_grain), already_loaded=already,
        top_unmatched=[(str(r[0]), float(r[1])) for r in top],
    )


def keep(upload: Upload, period: str, target_dir: Path | str) -> Path:
    """Write the month where a rebuild will find it again.

    An upload that only reaches the database is a month that disappears the
    next time the warehouse is rebuilt from source. Writing it beside the other
    committed months is what makes it part of the record.
    """
    target_dir = Path(target_dir)
    prefix = "pivot" if upload.kind == "workbook-pivot" else "long"
    target = target_dir / f"{prefix}_{period}.csv"
    write_period(upload, period, target)
    if prefix == "long":
        weaker = target_dir / f"pivot_{period}.csv"
        if weaker.exists():
            weaker.unlink()
    return target


def stage(uploaded_bytes: bytes, filename: str, directory: Path | str) -> Path:
    """Put an in-memory upload on disk so the readers can open it."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / Path(filename).name
    target.write_bytes(uploaded_bytes)
    return target
